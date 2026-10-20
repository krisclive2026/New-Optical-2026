from io import BytesIO
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill


def generate_report_excel(title, headers, rows):
    wb = Workbook()
    ws = wb.active
    ws.title = 'Report'
    ws.append([title])
    ws['A1'].font = Font(bold=True, size=13)
    ws.append([])
    ws.append(headers)
    for cell in ws[3]:
        cell.font = Font(bold=True, color='FFFFFF')
        cell.fill = PatternFill('solid', fgColor='0B3D6B')
    for r in rows:
        ws.append(list(r))
    for col in ws.columns:
        ws.column_dimensions[col[0].column_letter].width = min(max(len(str(c.value or '')) for c in col) + 3, 45)
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()
